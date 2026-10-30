import os
import re
import unicodedata

# 1. OCR Engine for Images & Scanned Documents
try:
    from rapidocr_onnxruntime import RapidOCR
    rapid_ocr = RapidOCR()
except Exception as e:
    rapid_ocr = None
    print(f"RapidOCR unavailable: {e}")

# 2. PDF parsers
try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import pypdfium2
except ImportError:
    pypdfium2 = None

try:
    from pdfminer.high_level import extract_text as pdfminer_extract_text
except ImportError:
    pdfminer_extract_text = None

# 3. Word DOCX parser
try:
    import docx
except ImportError:
    docx = None

# 4. Image processing
try:
    from PIL import Image
except ImportError:
    Image = None


def normalize_extracted_text(text: str) -> str:
    """
    Cleans raw text:
    - Normalizes unicode characters and ligatures (e.g. fi, fl, ffi)
    - Replaces odd bullets and dashes with standard characters
    - Normalizes inconsistent whitespace while preserving paragraph boundaries
    """
    if not text:
        return ""

    text = unicodedata.normalize("NFKD", text)

    replacements = {
        "\ufb00": "ff",
        "\ufb01": "fi",
        "\ufb02": "fl",
        "\ufb03": "ffi",
        "\ufb04": "ffl",
        "\u2022": "•",
        "\u2023": "•",
        "\u25e6": "•",
        "\u2043": "•",
        "\u2219": "•",
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\r\n": "\n",
        "\r": "\n",
    }
    for orig, repl in replacements.items():
        text = text.replace(orig, repl)

    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = [line.rstrip() for line in text.split("\n")]
    return "\n".join(lines).strip()


def extract_text_from_image(file_path: str) -> str:
    """
    Extracts text from images (PNG, JPG, JPEG, WEBP, BMP, TIFF) using RapidOCR.
    """
    if not rapid_ocr:
        return ""

    try:
        results, _ = rapid_ocr(file_path)
        if not results:
            return ""

        extracted_lines = []
        for item in results:
            # item format: [box, text, confidence]
            if len(item) >= 2 and item[1]:
                extracted_lines.append(item[1].strip())

        full_text = "\n".join(extracted_lines)
        return normalize_extracted_text(full_text)
    except Exception as e:
        print(f"Error during image OCR extraction on {file_path}: {e}")
        return ""


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extracts text from a PDF:
    1. Primary digital extraction with pdfplumber
    2. Fallback digital extraction with pypdfium2 / pdfminer
    3. Scanned PDF fallback: renders pages to images and runs OCR
    """
    text_chunks = []

    # Engine 1: pdfplumber (preserves layout, text blocks)
    if pdfplumber:
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text(layout=True) or page.extract_text()
                    if page_text and page_text.strip():
                        text_chunks.append(page_text.strip())
            if text_chunks:
                full_text = "\n\n".join(text_chunks)
                if len(full_text.strip()) > 60:
                    return normalize_extracted_text(full_text)
        except Exception as e:
            print(f"pdfplumber extraction notice: {e}")

    # Engine 2: pypdfium2 digital text extraction
    if pypdfium2:
        try:
            pdf = pypdfium2.PdfDocument(file_path)
            chunks = []
            for i in range(len(pdf)):
                page = pdf[i]
                textpage = page.get_textpage()
                page_text = textpage.get_text_range()
                if page_text and page_text.strip():
                    chunks.append(page_text.strip())
            if chunks:
                full_text = "\n\n".join(chunks)
                if len(full_text.strip()) > 60:
                    return normalize_extracted_text(full_text)
        except Exception as e:
            print(f"pypdfium2 extraction notice: {e}")

    # Engine 3: pdfminer
    if pdfminer_extract_text:
        try:
            mined_text = pdfminer_extract_text(file_path)
            if mined_text and len(mined_text.strip()) > 60:
                return normalize_extracted_text(mined_text)
        except Exception:
            pass

    # Engine 4: Scanned PDF OCR (if digital extraction returned little/no text)
    if rapid_ocr and pypdfium2:
        try:
            print(f"Running OCR on scanned PDF: {file_path}")
            pdf = pypdfium2.PdfDocument(file_path)
            ocr_chunks = []
            for i in range(min(len(pdf), 4)):  # max first 4 pages
                page = pdf[i]
                bitmap = page.render(scale=2.0)
                pil_image = bitmap.to_pil()
                # Run OCR directly on PIL image
                results, _ = rapid_ocr(pil_image)
                if results:
                    page_lines = [item[1].strip() for item in results if len(item) >= 2 and item[1]]
                    ocr_chunks.append("\n".join(page_lines))
            if ocr_chunks:
                full_text = "\n\n".join(ocr_chunks)
                if len(full_text.strip()) > 40:
                    return normalize_extracted_text(full_text)
        except Exception as e:
            print(f"Scanned PDF OCR failed: {e}")

    # Fallback to plain text read if file is not binary
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            raw_content = f.read()
            if raw_content and len(raw_content.strip()) > 30 and not raw_content.startswith("%PDF"):
                return normalize_extracted_text(raw_content)
    except Exception:
        pass

    return "\n\n".join(text_chunks) if text_chunks else ""


def extract_text_from_docx(file_path: str) -> str:
    """Extracts text from Microsoft Word (.docx) files."""
    if not docx:
        return ""
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))
        return normalize_extracted_text("\n".join(paragraphs))
    except Exception as e:
        print(f"Failed to extract DOCX: {e}")
        return ""


def extract_text_from_file(file_path: str) -> str:
    """
    Unified extraction entry point that detects file type and dispatches
    to the appropriate extractor (PDF, DOCX, Images, Text).
    """
    if not os.path.exists(file_path):
        return ""

    lower_path = file_path.lower()

    # Images
    if lower_path.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff")):
        return extract_text_from_image(file_path)

    # PDFs
    elif lower_path.endswith(".pdf"):
        return extract_text_from_pdf(file_path)

    # Word
    elif lower_path.endswith((".docx", ".doc")):
        return extract_text_from_docx(file_path)

    # Plain text
    elif lower_path.endswith((".txt", ".md", ".rtf")):
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return normalize_extracted_text(f.read())
        except Exception as e:
            print(f"Failed reading text file: {e}")
            return ""

    # Default fallback
    else:
        try:
            res = extract_text_from_pdf(file_path)
            if res:
                return res
        except Exception:
            pass
        try:
            return extract_text_from_image(file_path)
        except Exception:
            pass
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return normalize_extracted_text(f.read())
        except Exception:
            return ""
